#!/usr/bin/env python3
"""Lossless CSV/XLSX normalization. Never evaluates formulas or silently skips rows."""
import argparse,csv,hashlib,io,json
from pathlib import Path

def convert(path,sheet=None):
    path=Path(path)
    if path.suffix.lower()=='.csv':
        content=path.read_text(encoding='utf-8-sig')
        rows=list(csv.reader(io.StringIO(content),strict=True))
    elif path.suffix.lower()=='.xlsx':
        import openpyxl
        wb=openpyxl.load_workbook(path,read_only=True,data_only=False)
        if sheet is None and len(wb.sheetnames)!=1: raise ValueError('Multiple sheets: choose --sheet explicitly')
        ws=wb[sheet or wb.sheetnames[0]]
        rows=[]
        for row in ws.iter_rows():
            values=[]
            for cell in row:
                if cell.data_type=='f': raise ValueError(f'Formula at {cell.coordinate}: export calculated values first')
                if cell.data_type=='e': raise ValueError(f'Excel error at {cell.coordinate}')
                v=cell.value
                if v is None: v=''
                elif not isinstance(v,(str,int,float)): raise ValueError(f'Unsupported cell type at {cell.coordinate}: normalize explicitly')
                values.append(str(v))
            rows.append(values)
        wb.close()
    else: raise ValueError('Only .csv and .xlsx are supported')
    if len(rows)<2: raise ValueError('Header and at least one data row required')
    header=rows[0]
    if any(not h.strip() for h in header) or len(set(header))!=len(header): raise ValueError('Blank or duplicate header')
    for i,row in enumerate(rows[1:],1):
        if len(row)!=len(header): raise ValueError(f'Row {i}: wrong column count')
        if not any(v.strip() for v in row): raise ValueError(f'Row {i}: blank row; resolve explicitly')
    out=io.StringIO(newline='');csv.writer(out,lineterminator='\n').writerows(rows)
    content=out.getvalue()
    return content,{'source':str(path.name),'sheet':sheet,'expectedRows':len(rows)-1,'headers':header,'sourceSha256':hashlib.sha256(path.read_bytes()).hexdigest(),'csvSha256':hashlib.sha256(content.encode()).hexdigest(),'formulaPolicy':'reject','droppedRows':0}

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('input');ap.add_argument('--sheet');ap.add_argument('--output',required=True)
    a=ap.parse_args();content,manifest=convert(a.input,a.sheet);target=Path(a.output)
    if target.exists() or target.with_suffix('.manifest.json').exists(): raise ValueError('Output exists; choose a new path')
    target.parent.mkdir(parents=True,exist_ok=True);target.write_text(content,encoding='utf-8',newline='')
    target.with_suffix('.manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False))
